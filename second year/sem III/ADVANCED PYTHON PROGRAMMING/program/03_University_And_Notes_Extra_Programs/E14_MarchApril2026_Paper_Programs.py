# 1. March/April 2026 Q3(a): Extract numbers > 18 from x = [17, 19, 15, 3, 45] using filter()
x = [17, 19, 15, 3, 45]

def find(n):
    return n > 18

y = filter(find, x)
p = list(y)
print("Numbers > 18:", p)

# 2. March/April 2026 Q4(b): Read single cell B2 (Amit) from demo.xlsx and print on screen
from openpyxl import Workbook, load_workbook

# Creating demo.xlsx first so it can be read
wb_create = Workbook()
ws_create = wb_create.active
ws_create.append(["Sr. No.", "Name", "Place"])
ws_create.append([1, "Amit", "Pune"])
ws_create.append([2, "Rahul", "Solapur"])
ws_create.append([3, "Sachin", "Sangli"])
wb_create.save("demo.xlsx")

# Actual answer to read cell B2
wb = load_workbook('demo.xlsx')
ws = wb.active
cell_val = ws['B2']
print("Cell B2 value:", cell_val.value)
