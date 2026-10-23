# Program: Font Formatting On Name A
# Question: Write python program to perform font formatting on name A in sheet created by question no.
# 10.

# Setup: spreadsheet of Q.10 (Roll No. | Name) - skip if students.xlsx already exists
from openpyxl import Workbook
wb = Workbook()
ws = wb.active
ws.title = "Students"
ws.append(["Roll No.", "Name"])
ws.append([1, "A"])
ws.append([2, "B"])
ws.append([3, "C"])
wb.save("students.xlsx")

from openpyxl import load_workbook

from openpyxl.styles import Font
wb = load_workbook("students.xlsx")
ws = wb.active
ws["B2"].font = Font(name="Impact", bold=True, size=14, color="FF6347")   # name A is in B2
wb.save("students.xlsx")
print("Font formatting applied on name A")
