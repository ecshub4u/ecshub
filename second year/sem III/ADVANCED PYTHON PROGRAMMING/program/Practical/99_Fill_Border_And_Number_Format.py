# Program: Fill Border And Number Format
# Question: Write a program to apply fill, border and number format on cells.

from openpyxl import Workbook
from openpyxl.styles import PatternFill, Border, Side

wb = Workbook()
ws = wb.active
ws.append(["empid", "name"])
ws.append([101, "ABC"])
ws.append([102, "PQR"])

ws["B3"].fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

border_thick = Side(style="thick")
border_thin = Side(style="thin")
ws["F3"].border = Border(top=border_thick, left=border_thin, right=border_thin, bottom=border_thin)

ws["F6"] = 14342345.6779
ws["F6"].number_format = "#,##0.00"
ws["F7"] = "2024-01-01"
ws["F7"].number_format = "yyyy-mm-dd"

wb.save("styles.xlsx")
print("styles.xlsx created")
