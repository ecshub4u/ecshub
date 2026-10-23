# Program: Insert Two Rows Above Roll No. 3
# Question: Write python program to insert two rows above Roll No. 3.

# Setup: spreadsheet of Q.10 (Roll No. | Name) - skip if students.xlsx already exists
from openpyxl import Workbook
wb = Workbook()
ws = wb.active
ws.title = "Students"
ws.append(["Roll No.", "Name"])
ws.append([1, "A"])
ws.append([2, "B"])
ws.append([3, "C"])
wb.save("students.xlsx")

from openpyxl import load_workbook

wb = load_workbook("students.xlsx")
ws = wb.active
ws.insert_rows(idx=4, amount=2)     # Roll No. 3 is in row 4
wb.save("students.xlsx")
print("Two rows inserted above Roll No. 3")
