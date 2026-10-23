# Program: Insert And Delete Columns
# Question: Write a program to insert and delete columns in an Excel sheet.

from openpyxl import Workbook

wb = Workbook()
ws = wb.active
for row in [("Roll", "Name", "City"), (1, "A", "x"), (2, "B", "y"), (3, "C", "z")]:
    ws.append(row)

ws.insert_cols(idx=2, amount=1)      # insert column
print("After insert:")
for row in ws.iter_rows(values_only=True):
    print(row)

ws.delete_cols(idx=2, amount=1)      # delete column
print("After delete:")
for row in ws.iter_rows(values_only=True):
    print(row)

wb.save("columns.xlsx")
