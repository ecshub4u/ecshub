# Program: Align Text To Center
# Question: Write a program to align text to the center on spreadsheet created in Q.10.

# Setup: spreadsheet of Q.10 (Roll No. | Name) - skip if students.xlsx already exists
from openpyxl import Workbook
wb = Workbook()
ws = wb.active
ws.title = "Students"
ws.append(["Roll No.", "Name"])
ws.append([1, "A"])
ws.append([2, "B"])
ws.append([3, "C"])
wb.save("students.xlsx")

from openpyxl import load_workbook

from openpyxl.styles import Alignment
wb = load_workbook("students.xlsx")
ws = wb.active
for i in range(1, ws.max_row + 1):
    for j in range(1, ws.max_column + 1):
        ws.cell(row=i, column=j).alignment = Alignment(horizontal="center", vertical="center")
wb.save("students.xlsx")
print("Text aligned to center")
