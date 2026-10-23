# Program: Read And Display Excel On Screen
# Question: Write a python program to read and display the above excel on the screen.

# Setup: spreadsheet of Q.10 (Roll No. | Name) - skip if students.xlsx already exists
from openpyxl import Workbook
wb = Workbook()
ws = wb.active
ws.title = "Students"
ws.append(["Roll No.", "Name"])
ws.append([1, "A"])
ws.append([2, "B"])
ws.append([3, "C"])
wb.save("students.xlsx")

from openpyxl import load_workbook

wb = load_workbook("students.xlsx")
ws = wb.active
for i in range(1, ws.max_row + 1):
    for j in range(1, ws.max_column + 1):
        print(ws.cell(row=i, column=j).value, end="\t")
    print()
