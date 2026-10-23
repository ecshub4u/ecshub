# Program: Add Formulas SUM And AVERAGE
# Question: Write a program to add formulas (sum and average) in an Excel sheet.

from openpyxl import Workbook

wb = Workbook()
ws = wb.active
ws.append(["PHY.", "CHEM.", "MATHS", "Total", "Avg."])
ws.append([70, 77, 90])
ws["D2"] = "=SUM(A2,B2,C2)"
ws["E2"] = "=AVERAGE(A2,B2,C2)"
wb.save("formula.xlsx")
print("formula.xlsx created (open in Excel to see Total and Avg)")
