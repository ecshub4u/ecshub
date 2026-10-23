# Program: Conditional Formatting
# Question: Write a program to demonstrate conditional formatting on spreadsheet.

from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, FormulaRule

wb = Workbook()
ws = wb.active

data = [
    [10, 20, 30, 40],
    [15, 25, 35, 45],
    [50, 60, 70, 80],
    [5, 10, 15, 20],
]
for row in data:
    ws.append(row)

# 1) CellIsRule: highlight cells greater than 40 in red
red_fill = PatternFill(start_color="FFEE1110", end_color="FFEE1110", fill_type="solid")
ws.conditional_formatting.add("A1:D4", CellIsRule(operator="greaterThan", formula=["40"], fill=red_fill))

# 2) ColorScaleRule: 2-colour scale between 10 and 80
ws.conditional_formatting.add("A1:D4", ColorScaleRule(start_type="num", start_value=10, start_color="FFFFEDA0",
                                                      end_type="num", end_value=80, end_color="FF00BFFF"))

# 3) FormulaRule: highlight even numbers in blue
blue_fill = PatternFill(start_color="FF0000FF", end_color="FF0000FF", fill_type="solid")
ws.conditional_formatting.add("A1:D4", FormulaRule(formula=["MOD(A1,2)=0"], fill=blue_fill))

wb.save("conditional.xlsx")
print("conditional.xlsx created")
