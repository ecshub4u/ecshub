# Program: Replace Name C With D
# Question: Write a program to replace the name C as D.

# Setup: spreadsheet of Q.10 (Roll No. | Name) - skip if students.xlsx already exists
from openpyxl import Workbook
wb = Workbook()
ws = wb.active
ws.title = "Students"
ws.append(["Roll No.", "Name"])
ws.append([1, "A"])
ws.append([2, "B"])
ws.append([3, "C"])
wb.save("students.xlsx")

from openpyxl import load_workbook

wb = load_workbook("students.xlsx")
ws = wb.active
ws["B4"] = "D"                 # name C is in cell B4
wb.save("students.xlsx")
print("Name C replaced with D")
