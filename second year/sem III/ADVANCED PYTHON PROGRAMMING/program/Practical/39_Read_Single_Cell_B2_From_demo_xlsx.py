# Program: Read Single Cell B2 From demo.xlsx
# Question: Consider a spreadsheet named demo.xlsx (Sr. No., Name, Place: 1 Amit Pune, 2 Rahul
# Solapur, 3 Sachin Sangli). Write a python program to read single cell B2 (Amit) and print on screen.

# Setup: create demo.xlsx (in the exam this file already exists)
from openpyxl import Workbook
wb0 = Workbook()
ws0 = wb0.active
for row in [("Sr. No.", "Name", "Place"), (1, "Amit", "Pune"), (2, "Rahul", "Solapur"), (3, "Sachin", "Sangli")]:
    ws0.append(row)
wb0.save("demo.xlsx")

# demo.xlsx must be in the same folder (0_Sample_Files/create_sample_files.py makes it)
from openpyxl import load_workbook

wb = load_workbook("demo.xlsx")
ws = wb.active
x = ws["B2"]
y = ws.cell(row=2, column=2)
print(x.value)
print(y.value)
