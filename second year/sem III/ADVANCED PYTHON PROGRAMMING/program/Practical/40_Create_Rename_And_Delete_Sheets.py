# Program: Create Rename And Delete Sheets
# Question: Write a program to create a new sheet in the workbook, rename it, list sheet names and
# delete that sheet. (How can you manage - add, remove, rename - sheets using openpyxl?)

from openpyxl import Workbook

wb = Workbook()
ws = wb.active
ws.title = "MySheet"                         # rename a sheet

new_sheet = wb.create_sheet("New Title", 0)  # add a new sheet (index optional)
print(wb.sheetnames)                         # list all sheet names

new_sheet.title = "Renamed"                  # rename
print(wb.sheetnames)

del wb["Renamed"]                            # remove a sheet  (or wb.remove(sheet_object))
print(wb.sheetnames)

for sheet in wb:                             # loop through all sheets
    print(sheet.title)

wb.save("sheets.xlsx")
