# Program: Update Cell And Append Row
# Question: Write a program to create a workbook, add values to cells, update a cell and append a row.

from openpyxl import Workbook, load_workbook

wb = Workbook()
ws = wb.active
ws.title = "MySheet"
ws["A1"] = "emp id"
ws["B1"] = "emp name"
ws["A2"] = 101
ws["B2"] = "ABC"
wb.save("myworkbook.xlsx")

wb = load_workbook("myworkbook.xlsx")
ws = wb.active
ws["B2"] = "MNO"                     # updating cell value
r = (104, "NNN")
ws.append(r)                         # adding new row
wb.save("myworkbook.xlsx")
print("done")
