# Program: Freeze Columns
# Question: Write python program to freeze columns of spreadsheet created by question no. 10.

# Setup: spreadsheet of Q.10 (Roll No. | Name) - skip if students.xlsx already exists
from openpyxl import Workbook
wb = Workbook()
ws = wb.active
ws.title = "Students"
ws.append(["Roll No.", "Name"])
ws.append([1, "A"])
ws.append([2, "B"])
ws.append([3, "C"])
wb.save("students.xlsx")

from openpyxl import load_workbook

wb = load_workbook("students.xlsx")
ws = wb.active
ws.freeze_panes = "B1"              # freezes column A
wb.save("students.xlsx")
print("Column A frozen")
