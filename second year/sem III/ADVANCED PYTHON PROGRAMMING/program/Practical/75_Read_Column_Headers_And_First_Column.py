# Program: Read Column Headers And First Column
# Question: Write a program to read column headers and all values of the first column of an Excel
# sheet.

from openpyxl import Workbook, load_workbook

wb = Workbook()                      # create sample file
ws = wb.active
ws.append(["emp id", "emp name"])
ws.append([101, "ABC"])
ws.append([102, "PQR"])
ws.append([103, "XYZ"])
wb.save("new.xlsx")

wb = load_workbook("new.xlsx")
ws = wb.active
print("Column headers:")
for i in range(1, ws.max_column + 1):
    print(ws.cell(row=1, column=i).value)

print("First column:")
for i in range(1, ws.max_row + 1):
    print(ws.cell(row=i, column=1).value)
