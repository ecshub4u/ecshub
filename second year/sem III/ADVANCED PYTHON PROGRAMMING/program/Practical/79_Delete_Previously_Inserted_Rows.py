# Program: Delete Previously Inserted Rows
# Question: Write python program to delete previously inserted rows.

# Setup: spreadsheet of Q.10 (Roll No. | Name) - skip if students.xlsx already exists
from openpyxl import Workbook
wb = Workbook()
ws = wb.active
ws.title = "Students"
ws.append(["Roll No.", "Name"])
ws.append([1, "A"])
ws.append([2, "B"])
ws.append([3, "C"])
wb.save("students.xlsx")

# Setup: the two rows inserted by the previous program
from openpyxl import load_workbook
wb = load_workbook("students.xlsx")
ws = wb.active
ws.insert_rows(idx=4, amount=2)
wb.save("students.xlsx")

from openpyxl import load_workbook

wb = load_workbook("students.xlsx")
ws = wb.active
ws.delete_rows(idx=4, amount=2)     # delete the two inserted rows
wb.save("students.xlsx")
print("Inserted rows deleted")
