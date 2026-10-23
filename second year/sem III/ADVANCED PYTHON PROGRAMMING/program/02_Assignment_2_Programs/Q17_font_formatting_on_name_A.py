# Q17. Perform Font formatting on Name 'A' (Cell B2) using Font class (From Professor's Notes)
from openpyxl import load_workbook
from openpyxl.styles import Font

wb = load_workbook("student.xlsx")
ws = wb.active
ws['B2'].font = Font(name="Impact", bold=True, size=14, color="FF6347")
wb.save("student.xlsx")
print("Font formatting applied on Name A")
