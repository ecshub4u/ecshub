# Q24. Align text to the center on spreadsheet created in Q10 using Alignment class (From Professor's Notes)
from openpyxl import load_workbook
from openpyxl.styles import Alignment

wb = load_workbook("student.xlsx")
ws = wb.active
ws['A2'].alignment = Alignment(horizontal='center', vertical='center')
ws['B2'].alignment = Alignment(horizontal='center', vertical='center')
wb.save("student.xlsx")
print("Text aligned to center successfully")
