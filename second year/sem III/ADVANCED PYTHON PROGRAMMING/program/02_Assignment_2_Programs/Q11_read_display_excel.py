# Q11. Read and display Excel spreadsheet on screen (From Professor's Notes)
from openpyxl import load_workbook

wb = load_workbook("student.xlsx")
ws = wb.active

for i in range(1, ws.max_row + 1):
    for j in range(1, ws.max_column + 1):
        print(ws.cell(row=i, column=j).value, end="\t")
    print()
