# Q12. Replace the name C as D in cell B4 (From Professor's Notes)
from openpyxl import load_workbook

wb = load_workbook("student.xlsx")
ws = wb.active
ws['B4'] = "D"  # or ws.cell(row=4, column=2).value = "D"
wb.save("student.xlsx")
print("Replaced C with D successfully")
