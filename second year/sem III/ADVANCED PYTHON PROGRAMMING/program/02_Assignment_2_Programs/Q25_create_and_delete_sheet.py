# Q25. Create a new sheet in workbook as well as delete that sheet (From Professor's Notes)
from openpyxl import load_workbook

wb = load_workbook("student.xlsx")

# Create a new sheet
new_sheet = wb.create_sheet("NewSheet", 0)
print("Sheets after creation:", wb.sheetnames)

# Delete that sheet
del wb["NewSheet"]
print("Sheets after deletion:", wb.sheetnames)

wb.save("student.xlsx")
